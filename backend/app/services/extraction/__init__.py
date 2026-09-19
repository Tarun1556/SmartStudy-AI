from typing import List, Optional
import re


def extract_pdf_text(file_path: str) -> str:
    try:
        import fitz
        doc = fitz.open(file_path)
        parts = []
        for page in doc:
            parts.append(page.get_text())
        return "\n".join(parts)
    except Exception:
        return ""


def extract_pptx_text(file_path: str) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = "".join(run.text for run in para.runs)
                        if t.strip():
                            parts.append(t.strip())
                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if cells:
                            parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception:
        return ""


def extract_transcript_segments(file_path: str) -> List[dict]:
    ext = file_path.lower().rsplit(".", 1)[-1] if "." in file_path else ""
    if ext in ("pdf",):
        text = extract_pdf_text(file_path)
        return _text_to_segments(text, source_type="pdf")
    if ext in ("pptx", "ppt"):
        text = extract_pptx_text(file_path)
        return _text_to_segments(text, source_type="pptx")
    if ext in ("txt", "md", "rtf"):
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()
        except Exception:
            text = ""
        return _text_to_segments(text, source_type="text")
    return []


def _text_to_segments(text: str, source_type: str = "text") -> List[dict]:
    if not text:
        return []
    sentences = re.split(r'(?<=[.!?])\s+|\n+', text)
    segments = []
    buffer = ""
    idx = 0
    for s in sentences:
        s = s.strip()
        if not s:
            continue
        buffer = (buffer + " " + s).strip() if buffer else s
        if len(buffer) >= 300:
            segments.append({
                "segment_index": idx,
                "text": buffer,
                "start_time": None,
                "end_time": None,
                "source_type": source_type,
            })
            idx += 1
            buffer = ""
    if buffer:
        segments.append({
            "segment_index": idx,
            "text": buffer,
            "start_time": None,
            "end_time": None,
            "source_type": source_type,
        })
    return segments
